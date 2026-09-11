from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
from pathlib import Path
from xml.sax.saxutils import escape

from src.html_utils import clean_html
from src.pdf_generator import (
    organizar_criterios,
    organizar_contexto,
    organizar_requerimientos,
    obtener_prototipos
)

COLOR_GSE = RGBColor(0x18, 0xE0, 0xC4)
COLOR_LIGHTGREY = RGBColor(0xF2, 0xF2, 0xF2)
COLOR_BLACK = RGBColor(0x00, 0x00, 0x00)


def set_cell_shading(cell, color_hex):
    shading_elm = parse_xml(
        f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
    )
    cell._tc.get_or_add_tcPr().append(shading_elm)


def set_cell_borders(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '  <w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '  <w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '  <w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(tcBorders)


def set_cell_vertical_alignment(cell, align="center"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    val_map = {
        "top": "top",
        "center": "center",
        "bottom": "bottom"
    }
    vAlign = parse_xml(
        f'<w:vAlign {nsdecls("w")} w:val="{val_map.get(align, "center")}"/>'
    )
    tcPr.append(vAlign)


def set_paragraph_spacing(paragraph, before=0, after=0):
    pPr = paragraph._p.get_or_add_pPr()
    spacing = parse_xml(
        f'<w:spacing {nsdecls("w")} w:before="{before}" w:after="{after}"/>'
    )
    pPr.append(spacing)


def agregar_celda_contenido(cell, texto, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, font_size=10):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    run = p.add_run(str(texto))
    run.font.size = Pt(font_size)
    run.font.name = "Calibri"
    run.bold = bold
    set_paragraph_spacing(p, before=60, after=60)
    return p


def agregar_celda_html(cell, html_content, align=WD_ALIGN_PARAGRAPH.LEFT, font_size=10):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    set_paragraph_spacing(p, before=60, after=60)

    if not html_content:
        return p

    from bs4 import BeautifulSoup
    soup = BeautifulSoup(html_content, "html.parser")
    texto = soup.get_text(separator="\n")

    lines = [line.strip() for line in texto.split("\n") if line.strip()]

    for i, line in enumerate(lines):
        run = p.add_run(line)
        run.font.size = Pt(font_size)
        run.font.name = "Calibri"
        if i < len(lines) - 1:
            p.add_run("\n")

    return p


def crear_tabla(doc, num_rows, num_cols, col_widths_cm):
    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for row in table.rows:
        for cell in row.cells:
            set_cell_borders(cell)

    for i, width in enumerate(col_widths_cm):
        for row in table.rows:
            row.cells[i].width = Cm(width)

    return table


def titulo_seccion(doc, texto):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    cell = table.cell(0, 0)
    set_cell_shading(cell, "18E0C4")
    set_cell_borders(cell)

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=30, after=30)
    run = p.add_run(texto)
    run.bold = True
    run.font.size = Pt(11)
    run.font.name = "Calibri"


def generate_docx(
    work_item,
    total_historias_sprint,
    datos_requerimiento
):
    output_folder = Path("output")
    output_folder.mkdir(exist_ok=True)

    output_file = str(
        output_folder / f"HU_{work_item['id']}.docx"
    )

    doc = Document()

    section = doc.sections[0]
    section.top_margin = Cm(3)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

    logo_path = "src/assets/gse.png"
    if Path(logo_path).exists():
        header = section.header
        header.is_linked_to_previous = False

        # Crear tabla principal de encabezado (1 fila x 3 columnas)
        header_table = header.add_table(rows=1, cols=3, width=Cm(18))
        header_table.alignment = WD_TABLE_ALIGNMENT.CENTER

        # Columna 1: Logo (4cm)
        cell_logo = header_table.cell(0, 0)
        cell_logo.width = Cm(4)
        set_cell_borders(cell_logo)
        set_cell_vertical_alignment(cell_logo, "center")
        p_logo = cell_logo.paragraphs[0]
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_logo = p_logo.add_run()
        run_logo.add_picture(logo_path, width=Cm(3.5), height=Cm(2.5))

        # Columna 2: Título (8.5cm)
        cell_titulo = header_table.cell(0, 1)
        cell_titulo.width = Cm(8.5)
        set_cell_borders(cell_titulo)
        set_cell_vertical_alignment(cell_titulo, "center")
        p_titulo = cell_titulo.paragraphs[0]
        p_titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(p_titulo, before=30, after=30)
        run_titulo = p_titulo.add_run("HISTORIAS DE USUARIO")
        run_titulo.bold = True
        run_titulo.font.size = Pt(14)
        run_titulo.font.name = "Calibri"

        p_subtitulo = cell_titulo.add_paragraph()
        p_subtitulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_subtitulo = p_subtitulo.add_run("Desarrollo")
        run_subtitulo.font.size = Pt(11)
        run_subtitulo.font.name = "Calibri"

        # Columna 3: Metadata (5.5cm)
        cell_info = header_table.cell(0, 2)
        cell_info.width = Cm(5.5)
        set_cell_borders(cell_info)

        # Crear tabla interna para metadata
        info_table = cell_info.add_table(rows=4, cols=2)
        info_table.allow_autofit = False

        # Establecer anchos de columna: 2.5cm + 3cm = 5.5cm
        for row in info_table.rows:
            row.cells[0].width = Cm(2.5)
            row.cells[1].width = Cm(3)

        info_data = [
            ("Código", "PTI-DS-FR-84"),
            ("Versión", "6"),
            ("Implementación", "01/08/2025"),
            ("Clasificación de\nla información", "Uso Interno")
        ]

        for i, (label, value) in enumerate(info_data):
            cell_label = info_table.cell(i, 0)
            cell_value = info_table.cell(i, 1)
            set_cell_shading(cell_label, "F2F2F2")
            set_cell_borders(cell_label)
            set_cell_borders(cell_value)
            agregar_celda_contenido(cell_label, label, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=9)
            agregar_celda_contenido(cell_value, value, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=9)

        # Agregar espacio entre encabezado y contenido
        spacer_para = header.add_paragraph()
        set_paragraph_spacing(spacer_para, before=0, after=200)

    # ==================================================
    # FECHA Y PROYECTO
    # ==================================================
    fecha = work_item["fields"].get("System.CreatedDate", "")
    dia = ""
    mes = ""
    año = ""

    if fecha:
        fecha_limpia = fecha[:10]
        año, mes, dia = fecha_limpia.split("-")

    table_fecha = doc.add_table(rows=2, cols=4)
    table_fecha.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Establecer anchos de columna: 2cm, 2cm, 3cm, 11cm
    for row in table_fecha.rows:
        row.cells[0].width = Cm(2)
        row.cells[1].width = Cm(2)
        row.cells[2].width = Cm(3)
        row.cells[3].width = Cm(11)

    cell_fecha_label = table_fecha.cell(0, 0).merge(table_fecha.cell(0, 2))
    agregar_celda_contenido(cell_fecha_label, "FECHA", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_shading(cell_fecha_label, "18E0C4")

    cell_proyecto = table_fecha.cell(0, 3)
    agregar_celda_contenido(cell_proyecto, "PROYECTO", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_shading(cell_proyecto, "18E0C4")

    cell_dia_label = table_fecha.cell(1, 0)
    agregar_celda_contenido(cell_dia_label, "DIA", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_shading(cell_dia_label, "F2F2F2")

    cell_mes_label = table_fecha.cell(1, 1)
    agregar_celda_contenido(cell_mes_label, "MES", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_shading(cell_mes_label, "F2F2F2")

    cell_anio_label = table_fecha.cell(1, 2)
    agregar_celda_contenido(cell_anio_label, "AÑO", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_shading(cell_anio_label, "F2F2F2")

    cell_proyecto_valor = table_fecha.cell(1, 3)
    cell_proyecto_valor.merge(table_fecha.cell(1, 3))
    agregar_celda_contenido(
        cell_proyecto_valor,
        work_item["fields"].get("System.TeamProject", ""),
        align=WD_ALIGN_PARAGRAPH.CENTER
    )

    cell_dia = table_fecha.cell(1, 0)
    cell_dia.text = ""
    agregar_celda_contenido(cell_dia, dia, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_shading(cell_dia, "F2F2F2")

    cell_mes = table_fecha.cell(1, 1)
    cell_mes.text = ""
    agregar_celda_contenido(cell_mes, mes, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_shading(cell_mes, "F2F2F2")

    cell_anio = table_fecha.cell(1, 2)
    cell_anio.text = ""
    agregar_celda_contenido(cell_anio, año, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_shading(cell_anio, "F2F2F2")

    for row in table_fecha.rows:
        for cell in row.cells:
            set_cell_borders(cell)

    doc.add_paragraph()

    # ==================================================
    # SECCIÓN 1 Y 2
    # ==================================================
    req_info = None
    if datos_requerimiento and isinstance(datos_requerimiento, dict):
        req_info = datos_requerimiento.get("predecesor")

    if not req_info:
        req_info = {
            "id_requerimiento": "N/A",
            "nombre_requerimiento": "No asignado",
            "descripcion": "No se encontró un requerimiento predecesor vinculado."
        }

    nombre_requerimiento = req_info["nombre_requerimiento"]
    id_requerimiento = req_info["id_requerimiento"]

    table12 = doc.add_table(rows=2, cols=2)
    table12.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Establecer anchos de columna: 13cm, 5cm
    for row in table12.rows:
        row.cells[0].width = Cm(13)
        row.cells[1].width = Cm(5)

    cell_1_label = table12.cell(0, 0)
    agregar_celda_contenido(cell_1_label, "1. Nombre del requerimiento", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_shading(cell_1_label, "18E0C4")

    cell_2_label = table12.cell(0, 1)
    agregar_celda_contenido(cell_2_label, "2. ID requerimiento", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_shading(cell_2_label, "18E0C4")

    cell_1_valor = table12.cell(1, 0)
    agregar_celda_contenido(cell_1_valor, nombre_requerimiento, align=WD_ALIGN_PARAGRAPH.CENTER)

    cell_2_valor = table12.cell(1, 1)
    agregar_celda_contenido(cell_2_valor, str(id_requerimiento), align=WD_ALIGN_PARAGRAPH.CENTER)

    for row in table12.rows:
        for cell in row.cells:
            set_cell_borders(cell)

    doc.add_paragraph()

    # ==================================================
    # SECCIÓN 3
    # ==================================================
    titulo_seccion(doc, "3. DESCRIPCIÓN DEL REQUERIMIENTO")

    descripcion = req_info["descripcion"]

    table3 = doc.add_table(rows=1, cols=1)
    table3.alignment = WD_TABLE_ALIGNMENT.CENTER

    cell3 = table3.cell(0, 0)
    agregar_celda_html(cell3, descripcion, font_size=10)
    set_cell_vertical_alignment(cell3, "top")

    for row in table3.rows:
        for cell in row.cells:
            set_cell_borders(cell)

    doc.add_paragraph()

    # ==================================================
    # SECCIÓN 4 Y 5
    # ==================================================
    table45 = doc.add_table(rows=2, cols=2)
    table45.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Establecer anchos de columna: 9cm, 9cm
    for row in table45.rows:
        row.cells[0].width = Cm(9)
        row.cells[1].width = Cm(9)

    cell_4_label = table45.cell(0, 0)
    agregar_celda_contenido(cell_4_label, "4. Tipo de requerimiento (Negocio, Técnico, Soporte)", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_shading(cell_4_label, "18E0C4")

    cell_5_label = table45.cell(0, 1)
    agregar_celda_contenido(cell_5_label, "5. Total historias de usuario", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_shading(cell_5_label, "18E0C4")

    cell_4_valor = table45.cell(1, 0)
    agregar_celda_contenido(cell_4_valor, "Técnico", align=WD_ALIGN_PARAGRAPH.CENTER)

    cell_5_valor = table45.cell(1, 1)
    agregar_celda_contenido(cell_5_valor, str(total_historias_sprint), align=WD_ALIGN_PARAGRAPH.CENTER)

    for row in table45.rows:
        for cell in row.cells:
            set_cell_borders(cell)

    doc.add_paragraph()

    # ==================================================
    # SECCIÓN 6 Y 7
    # ==================================================
    prioridad = work_item["fields"].get("Microsoft.VSTS.Common.Priority", "")
    prioridad_texto = {
        1: "Baja",
        2: "Media",
        3: "Alta",
        4: "Alta"
    }.get(prioridad, "N/A")

    table67 = doc.add_table(rows=2, cols=2)
    table67.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Establecer anchos de columna: 9cm, 9cm
    for row in table67.rows:
        row.cells[0].width = Cm(9)
        row.cells[1].width = Cm(9)

    cell_6_label = table67.cell(0, 0)
    agregar_celda_contenido(cell_6_label, "6. Prioridad (Alta, Media o Baja)", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_shading(cell_6_label, "18E0C4")

    cell_7_label = table67.cell(0, 1)
    agregar_celda_contenido(cell_7_label, "7. ID Historia", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_shading(cell_7_label, "18E0C4")

    cell_6_valor = table67.cell(1, 0)
    agregar_celda_contenido(cell_6_valor, prioridad_texto, align=WD_ALIGN_PARAGRAPH.CENTER)

    cell_7_valor = table67.cell(1, 1)
    agregar_celda_contenido(cell_7_valor, str(work_item["id"]), align=WD_ALIGN_PARAGRAPH.CENTER)

    for row in table67.rows:
        for cell in row.cells:
            set_cell_borders(cell)

    doc.add_paragraph()

    # ==================================================
    # SECCIÓN 8 Y 9
    # ==================================================
    complejidad = work_item["fields"].get("Microsoft.VSTS.Common.Risk", "")
    complejidad_texto = {
        "1 - High": "Alta",
        "2 - Medium": "Media",
        "3 - Low": "Baja"
    }.get(complejidad, "N/A")

    estimacion = work_item["fields"].get("Microsoft.VSTS.Scheduling.StoryPoints", "")
    if estimacion not in ("", None):
        estimacion = int(estimacion)
    else:
        estimacion = "N/A"

    table89 = doc.add_table(rows=2, cols=2)
    table89.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Establecer anchos de columna: 9cm, 9cm
    for row in table89.rows:
        row.cells[0].width = Cm(9)
        row.cells[1].width = Cm(9)

    cell_8_label = table89.cell(0, 0)
    agregar_celda_contenido(cell_8_label, "8. Complejidad (Alta, Media o Baja)", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_shading(cell_8_label, "18E0C4")

    cell_9_label = table89.cell(0, 1)
    agregar_celda_contenido(cell_9_label, "9. Estimación", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_shading(cell_9_label, "18E0C4")

    cell_8_valor = table89.cell(1, 0)
    agregar_celda_contenido(cell_8_valor, complejidad_texto, align=WD_ALIGN_PARAGRAPH.CENTER)

    cell_9_valor = table89.cell(1, 1)
    agregar_celda_contenido(cell_9_valor, str(estimacion), align=WD_ALIGN_PARAGRAPH.CENTER)

    for row in table89.rows:
        for cell in row.cells:
            set_cell_borders(cell)

    doc.add_paragraph()

    # ==================================================
    # SECCIÓN 10
    # ==================================================
    requerimiento = work_item["fields"].get("System.Title", "")

    table10 = doc.add_table(rows=2, cols=1)
    table10.alignment = WD_TABLE_ALIGNMENT.CENTER

    cell_10_label = table10.cell(0, 0)
    agregar_celda_contenido(cell_10_label, "10. Nombre Historia de Usuario", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_shading(cell_10_label, "18E0C4")

    cell_10_valor = table10.cell(1, 0)
    agregar_celda_contenido(cell_10_valor, requerimiento, align=WD_ALIGN_PARAGRAPH.CENTER)

    for row in table10.rows:
        for cell in row.cells:
            set_cell_borders(cell)

    doc.add_paragraph()

    # ==================================================
    # SECCIÓN 11
    # ==================================================
    titulo_seccion(doc, "11. Descripción de la historia de usuario")

    como = clean_html(work_item["fields"].get("Custom.Como", ""))
    quiero = clean_html(work_item["fields"].get("System.Description", ""))
    para = clean_html(work_item["fields"].get("Custom.Para", ""))
    contexto = organizar_contexto(work_item["fields"].get("Custom.Contexto", ""))

    table11 = doc.add_table(rows=1, cols=1)
    table11.alignment = WD_TABLE_ALIGNMENT.CENTER

    cell11 = table11.cell(0, 0)
    cell11.text = ""

    p_como_label = cell11.paragraphs[0]
    run_como_label = p_como_label.add_run("Como:")
    run_como_label.bold = True
    run_como_label.font.size = Pt(10)
    run_como_label.font.name = "Calibri"

    p_como = cell11.add_paragraph()
    run_como = p_como.add_run(como)
    run_como.font.size = Pt(10)
    run_como.font.name = "Calibri"

    cell11.add_paragraph()

    p_quiero_label = cell11.add_paragraph()
    run_quiero_label = p_quiero_label.add_run("Quiero:")
    run_quiero_label.bold = True
    run_quiero_label.font.size = Pt(10)
    run_quiero_label.font.name = "Calibri"

    p_quiero = cell11.add_paragraph()
    run_quiero = p_quiero.add_run(quiero)
    run_quiero.font.size = Pt(10)
    run_quiero.font.name = "Calibri"

    cell11.add_paragraph()

    p_para_label = cell11.add_paragraph()
    run_para_label = p_para_label.add_run("Con la finalidad de:")
    run_para_label.bold = True
    run_para_label.font.size = Pt(10)
    run_para_label.font.name = "Calibri"

    p_para = cell11.add_paragraph()
    run_para = p_para.add_run(para)
    run_para.font.size = Pt(10)
    run_para.font.name = "Calibri"

    cell11.add_paragraph()

    p_contexto_label = cell11.add_paragraph()
    run_contexto_label = p_contexto_label.add_run("Cuando:")
    run_contexto_label.bold = True
    run_contexto_label.font.size = Pt(10)
    run_contexto_label.font.name = "Calibri"

    for line in contexto.split("<br/>"):
        p_contexto = cell11.add_paragraph()
        run_contexto = p_contexto.add_run(line)
        run_contexto.font.size = Pt(10)
        run_contexto.font.name = "Calibri"

    for row in table11.rows:
        for cell in row.cells:
            set_cell_borders(cell)

    doc.add_paragraph()

    # ==================================================
    # SECCIÓN 12
    # ==================================================
    titulo_seccion(doc, "12. Fuera de alcance")

    requerimientos = work_item["fields"].get("Custom.Requerimientos", "")
    requerimientos_numerados = organizar_requerimientos(requerimientos)

    table12_fuera = doc.add_table(rows=1, cols=1)
    table12_fuera.alignment = WD_TABLE_ALIGNMENT.CENTER

    cell12f = table12_fuera.cell(0, 0)
    cell12f.text = ""

    if requerimientos_numerados:
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(requerimientos_numerados, "html.parser")
        texto = soup.get_text(separator="\n")
        lines = [line.strip() for line in texto.split("\n") if line.strip()]
        for i, line in enumerate(lines):
            if i == 0:
                p = cell12f.paragraphs[0]
            else:
                p = cell12f.add_paragraph()
            run = p.add_run(line)
            run.font.size = Pt(10)
            run.font.name = "Calibri"
    else:
        p = cell12f.paragraphs[0]
        run = p.add_run("N/A.")
        run.font.size = Pt(10)
        run.font.name = "Calibri"

    for row in table12_fuera.rows:
        for cell in row.cells:
            set_cell_borders(cell)

    doc.add_paragraph()

    # ==================================================
    # SECCIÓN 13
    # ==================================================
    titulo_seccion(doc, "13. Criterios de aceptación")

    criterios_html = work_item["fields"].get(
        "Microsoft.VSTS.Common.AcceptanceCriteria", ""
    )
    criterios_organizados = organizar_criterios(criterios_html)

    table13 = doc.add_table(rows=1, cols=1)
    table13.alignment = WD_TABLE_ALIGNMENT.CENTER

    cell13 = table13.cell(0, 0)
    cell13.text = ""

    if criterios_organizados:
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(criterios_organizados, "html.parser")
        texto = soup.get_text(separator="\n")
        lines = [line.strip() for line in texto.split("\n") if line.strip()]
        for i, line in enumerate(lines):
            if i == 0:
                p = cell13.paragraphs[0]
            else:
                p = cell13.add_paragraph()
            run = p.add_run(line)
            run.font.size = Pt(10)
            run.font.name = "Calibri"
    else:
        p = cell13.paragraphs[0]
        run = p.add_run("N/A.")
        run.font.size = Pt(10)
        run.font.name = "Calibri"

    for row in table13.rows:
        for cell in row.cells:
            set_cell_borders(cell)

    doc.add_paragraph()

    # ==================================================
    # SECCIÓN 14
    # ==================================================
    titulo_seccion(doc, "14. Criterios de seguridad")

    table14 = doc.add_table(rows=5, cols=2)
    table14.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Establecer anchos de columna: 7cm, 11cm
    for row in table14.rows:
        row.cells[0].width = Cm(7)
        row.cells[1].width = Cm(11)

    cell_criterio_header = table14.cell(0, 0)
    agregar_celda_contenido(cell_criterio_header, "Criterio", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_shading(cell_criterio_header, "F2F2F2")

    cell_detalle_header = table14.cell(0, 1)
    agregar_celda_contenido(cell_detalle_header, "Detalle", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_shading(cell_detalle_header, "F2F2F2")

    seguridad_data = [
        ("Mecanismo para validar proceso de Autenticación y Autorización",
         "• Acceso a la plataforma mediante roles 2FA para inicio de sesión (App de authenticator).\n• Envío de código OTP para la comprobación de correo electrónico en el registro token jwt librería de aws Cognito"),
        ("Cifrado de datos sensibles",
         "• Cifrado post-cuántico con el algoritmo de Kyber KEM"),
        ("Mecanismo de comprobación de integridad de los datos",
         "• AES (Advanced Encryption Standard).\n• SHA-256"),
        ("Inclusión de logs de trazabilidad",
         "• Almacenamiento de transacciones de blockchain en base de datos")
    ]

    for i, (criterio, detalle) in enumerate(seguridad_data, start=1):
        cell_c = table14.cell(i, 0)
        agregar_celda_contenido(cell_c, criterio, align=WD_ALIGN_PARAGRAPH.CENTER)

        cell_d = table14.cell(i, 1)
        cell_d.text = ""
        for j, line in enumerate(detalle.split("\n")):
            if j == 0:
                p = cell_d.paragraphs[0]
            else:
                p = cell_d.add_paragraph()
            run = p.add_run(line)
            run.font.size = Pt(10)
            run.font.name = "Calibri"

    for row in table14.rows:
        for cell in row.cells:
            set_cell_borders(cell)

    doc.add_paragraph()

    # ==================================================
    # SECCIÓN 15
    # ==================================================
    titulo_seccion(doc, "15. Interfaz gráfica")

    prototipo_html = work_item["fields"].get("Custom.PrototiposoWireframes", "")
    prototipos = obtener_prototipos(prototipo_html)

    table15 = doc.add_table(rows=1, cols=1)
    table15.alignment = WD_TABLE_ALIGNMENT.CENTER

    cell15 = table15.cell(0, 0)
    cell15.text = ""

    if not prototipos:
        p = cell15.paragraphs[0]
        run = p.add_run("N/A")
        run.font.size = Pt(10)
        run.font.name = "Calibri"
    else:
        for index, prototipo in enumerate(prototipos, start=1):
            p_desc = cell15.add_paragraph()
            run_desc = p_desc.add_run(f"{index}. {prototipo['descripcion']}")
            run_desc.bold = True
            run_desc.font.size = Pt(10)
            run_desc.font.name = "Calibri"

            p_url = cell15.add_paragraph()
            run_url = p_url.add_run(f"Ver prototipo en Figma: {prototipo['url']}")
            run_url.font.size = Pt(10)
            run_url.font.name = "Calibri"
            run_url.font.color.rgb = RGBColor(0x00, 0x56, 0xB3)

    for row in table15.rows:
        for cell in row.cells:
            set_cell_borders(cell)

    doc.add_paragraph()

    # ==================================================
    # SECCIÓN 16
    # ==================================================
    titulo_seccion(doc, "16. ¿Depende de otras historias de usuario?")

    dependencias_list = []

    if datos_requerimiento and isinstance(datos_requerimiento, dict):
        lista_relacionados = datos_requerimiento.get("relacionado", [])

        for rel in lista_relacionados:
            id_rel = rel.get("id_relacionado", "")
            titulo_rel = rel.get("titulo", "")

            if titulo_rel:
                titulo_escapado = escape(titulo_rel)
                dependencias_list.append(
                    f"{id_rel} - {titulo_escapado}" if id_rel else titulo_escapado
                )

    if dependencias_list:
        dependencia_texto = "\n".join(dependencias_list)
    else:
        dependencia_texto = "N/A."

    table16 = doc.add_table(rows=1, cols=1)
    table16.alignment = WD_TABLE_ALIGNMENT.CENTER

    cell16 = table16.cell(0, 0)
    cell16.text = ""

    from bs4 import BeautifulSoup
    soup = BeautifulSoup(dependencia_texto, "html.parser")
    texto = soup.get_text(separator="\n")
    lines = [line.strip() for line in texto.split("\n") if line.strip()]

    for i, line in enumerate(lines):
        if i == 0:
            p = cell16.paragraphs[0]
        else:
            p = cell16.add_paragraph()
        run = p.add_run(line)
        run.font.size = Pt(10)
        run.font.name = "Calibri"

    for row in table16.rows:
        for cell in row.cells:
            set_cell_borders(cell)

    doc.add_paragraph()

    # ==================================================
    # SECCIÓN 17
    # ==================================================
    titulo_seccion(doc, "17. Anexos")

    table17 = doc.add_table(rows=1, cols=1)
    table17.alignment = WD_TABLE_ALIGNMENT.CENTER

    cell17 = table17.cell(0, 0)
    agregar_celda_contenido(cell17, "N/A", align=WD_ALIGN_PARAGRAPH.LEFT)

    for row in table17.rows:
        for cell in row.cells:
            set_cell_borders(cell)

    doc.add_paragraph()

    # ==================================================
    # ELABORADO POR
    # ==================================================
    fecha = work_item["fields"].get("System.CreatedDate", "")
    dia = ""
    mes = ""
    año = ""
    fecha_formateada = "N/A"

    if fecha:
        fecha_limpia = fecha[:10]
        año, mes, dia = fecha_limpia.split("-")
        fecha_formateada = f"{dia}/{mes}/{año}"

    table_elaborado = doc.add_table(rows=2, cols=4)
    table_elaborado.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers_elaborado = ["Elaborado por", "Fecha", "Creación/Modificación", "Versión"]
    for i, header in enumerate(headers_elaborado):
        cell = table_elaborado.cell(0, i)
        agregar_celda_contenido(cell, header, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(cell, "18E0C4")

    valores_elaborado = ["Linda Daniela Corchuelo Pachon", fecha_formateada, "Creación", "1.0"]
    for i, valor in enumerate(valores_elaborado):
        cell = table_elaborado.cell(1, i)
        agregar_celda_contenido(cell, valor, align=WD_ALIGN_PARAGRAPH.CENTER)

    for row in table_elaborado.rows:
        for cell in row.cells:
            set_cell_borders(cell)

    doc.add_paragraph()

    # ==================================================
    # APROBADO POR
    # ==================================================
    titulo_seccion(doc, "Aprobado por")

    table_aprobado = doc.add_table(rows=4, cols=5)
    table_aprobado.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers_aprobado = ["Nombre", "Cargo", "Área/Empresa", "Firma", "Fecha"]
    for i, header in enumerate(headers_aprobado):
        cell = table_aprobado.cell(0, i)
        agregar_celda_contenido(cell, header, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(cell, "F2F2F2")

    aprobadores = [
        ("Linda Daniela Corchuelo Pachon", "Analista de requerimientos Junior Nivel 3", "Tecnología", "", ""),
        ("Eduin Fabian Ordonez Parra", "Product Manager Specialist", "Tecnología", "", ""),
        ("Carlos Alberto Rodriguez Sanchez", "Líder Técnico Nivel 1", "Tecnología", "", "")
    ]

    for row_idx, (nombre, cargo, area, firma, fecha_aprob) in enumerate(aprobadores, start=1):
        cell_nombre = table_aprobado.cell(row_idx, 0)
        agregar_celda_contenido(cell_nombre, nombre)

        cell_cargo = table_aprobado.cell(row_idx, 1)
        agregar_celda_contenido(cell_cargo, cargo)

        cell_area = table_aprobado.cell(row_idx, 2)
        agregar_celda_contenido(cell_area, area, align=WD_ALIGN_PARAGRAPH.CENTER)

        cell_firma = table_aprobado.cell(row_idx, 3)
        agregar_celda_contenido(cell_firma, firma, align=WD_ALIGN_PARAGRAPH.CENTER)

        cell_fecha = table_aprobado.cell(row_idx, 4)
        agregar_celda_contenido(cell_fecha, fecha_aprob, align=WD_ALIGN_PARAGRAPH.CENTER)

    for row in table_aprobado.rows:
        for cell in row.cells:
            set_cell_borders(cell)

    # ==================================================
    # GUARDAR DOCUMENTO
    # ==================================================
    doc.save(output_file)
    print(f"\nDOCX generado correctamente: {output_file}")

    return output_file
